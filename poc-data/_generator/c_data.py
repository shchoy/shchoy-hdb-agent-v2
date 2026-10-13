"""Excel workbooks, CSV/JSON system exports, and the relational database export."""
import csv
import datetime as dt
import json
import random
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import D, addr, arrears_asof, iso, register, sname, tenancy, tenant, town_name, unit, vacancy_asof
from master import APP_TYPE_LABEL, CHECKLISTS, DEFECT_TAXONOMY, SEVERITY_SLA_DAYS, STAFF, TAX_BY_CODE, TOWNS, TRADES
from render import DATEF, MONEY, TITLE_FONT, xl_header, xl_props, xl_widths


def _title(ws, text, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="555555")


def _rows(ws, start, rows, fmts=None):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=start + i, column=j + 1, value=v)
            if fmts and fmts.get(j):
                c.number_format = fmts[j]
    return start + len(rows) - 1


# ------------------------------------------------------------------ XLSX builders
def x_vacancy(M, path):
    asof = D(2026, 8, 31)
    vac = vacancy_asof(M, asof)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _title(ws, "HOUSING & ESTATES BOARD - PROPERTIES & LAND GROUP", "MONTHLY VACANCY REPORT - AS AT 31 AUG 2026", 8)
    ws.merge_cells("A4:A5")
    ws["A4"] = "Town"
    ws.merge_cells("B4:C4")
    ws["B4"] = "Units"
    ws.merge_cells("D4:F4")
    ws["D4"] = "Non-tenanted"
    ws.merge_cells("G4:H4")
    ws["G4"] = "Rates"
    xl_header(ws, 5, ["Total", "Tenanted", "Vacant", "Under Tender", "Not Lettable", "Vacancy rate", "Occupancy"], col=2)
    for c in "ABDG":
        ws[f"{c}4"].font = Font(bold=True)
        ws[f"{c}4"].alignment = Alignment(horizontal="center")
        ws[f"{c}4"].fill = PatternFill("solid", fgColor="B4C6E7")
    r = 6
    for t in TOWNS:
        c = t["code"]
        ws.cell(r, 1, t["name"])
        ws.cell(r, 2, vac["total"][c])
        ws.cell(r, 4, vac["vacant"][c])
        ws.cell(r, 5, vac["tender"][c])
        ws.cell(r, 6, vac["not_lettable"][c])
        ws.cell(r, 3, f"=B{r}-D{r}-E{r}-F{r}")
        ws.cell(r, 7, f"=D{r}/(B{r}-F{r})").number_format = "0.0%"
        ws.cell(r, 8, f"=C{r}/(B{r}-F{r})").number_format = "0.0%"
        r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    for col in "BCDEF":
        ws[f"{col}{r}"] = f"=SUM({col}6:{col}{r - 1})"
        ws[f"{col}{r}"].font = Font(bold=True)
    ws[f"G{r}"] = f"=D{r}/(B{r}-F{r})"
    ws[f"H{r}"] = f"=C{r}/(B{r}-F{r})"
    ws[f"G{r}"].number_format = ws[f"H{r}"].number_format = "0.0%"
    total_v = sum(vac["vacant"].values())
    notes = ["Notes:",
             f"1. Total vacant units as at 31 Aug 2026: {total_v}.",
             "2. Vacant EXCLUDES units under tender and not-lettable units w.e.f. April 2026 report (PLG Mgmt Meeting 4/2026, item 3).",
             "3. Prima Dental (U-ST512-01-33) tender awarded Jul 2026, commencement 1 Oct 2026 - counted as tenanted (committed).",
             "4. Source: PLG Records extract 01-Sep-2026 07:00. Prepared by K. Goh."]
    for i, n in enumerate(notes):
        ws.cell(r + 2 + i, 1, n).font = Font(italic=True, bold=(i == 0))
    xl_widths(ws, [18, 10, 11, 10, 13, 13, 13, 12])
    ws2 = wb.create_sheet("Vacant Units")
    hdr = ["Unit ID", "Town", "Address", "Unit type", "Floor area (sqm)", "NC", "Vacant since", "Days vacant at 31-Aug"]
    xl_header(ws2, 1, hdr)
    vu = sorted([u for u in M["units"] if u["status"] == "Vacant" and u["vacant_since"] <= asof], key=lambda u: u["unit_id"])
    for i, u in enumerate(vu, 2):
        _rows(ws2, i, [[u["unit_id"], town_name(u["town"]), addr(u), u["unit_type"], u["floor_area_sqm"], u["nc_id"], u["vacant_since"], f"=DATE(2026,8,31)-G{i}"]],
              {6: DATEF})
    xl_widths(ws2, [16, 14, 52, 24, 14, 10, 13, 18])
    ws3 = wb.create_sheet("Under Tender")
    xl_header(ws3, 1, ["Unit ID", "Town", "Address", "Unit type", "Floor area (sqm)", "Since"])
    for i, u in enumerate(sorted([u for u in M["units"] if u["status"] == "Under Tender"], key=lambda u: u["unit_id"]), 2):
        _rows(ws3, i, [[u["unit_id"], town_name(u["town"]), addr(u), u["unit_type"], u["floor_area_sqm"], u["vacant_since"]]], {5: DATEF})
    xl_widths(ws3, [16, 14, 52, 24, 14, 13])
    xl_props(wb, "Monthly Vacancy Report Aug 2026", sname("S08"), D(2026, 9, 2))
    wb.save(path)


def x_arrears(M, path):
    asof = D(2026, 9, 30)
    arr = arrears_asof(M, asof)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    hdr = ["Tenancy ID", "Tenant ID", "Tenant name", "Unit ID", "Town", "Monthly rent", "Unpaid invoices", "Oldest due date", "Days overdue",
           "Arrears (incl GST)", "Months of rent", "Bucket", "> 3 months?"]
    xl_header(ws, 1, hdr)
    for i, r in enumerate(arr, 2):
        ta = tenancy(M, r["tenancy_id"])
        _rows(ws, i, [[r["tenancy_id"], r["tenant_id"], tenant(M, r["tenant_id"])["name"], r["unit_id"], town_name(unit(M, r["unit_id"])["town"]), ta["monthly_rent"],
                       r["invoices"], r["oldest_due"], r["days_overdue"], r["amount"], f"=J{i}/(F{i}*1.09)", r["bucket"], f'=IF(K{i}>3,"YES","")']],
              {5: MONEY, 7: DATEF, 9: MONEY, 10: "0.00"})
    n = len(arr) + 1
    xl_widths(ws, [15, 10, 34, 16, 14, 12, 10, 13, 10, 15, 10, 8, 10])
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Unpaid Invoices")
    xl_header(ws2, 1, ["Invoice ID", "Tenancy ID", "Period", "Due date", "Total (incl GST)"])
    invs = [i for i in M["invoices"] if i["due_date"] <= asof and i["status"] == "Unpaid"]
    for j, inv in enumerate(invs, 2):
        _rows(ws2, j, [[inv["invoice_id"], inv["tenancy_id"], inv["period"], inv["due_date"], inv["total_amount"]]], {3: DATEF, 4: MONEY})
    xl_widths(ws2, [18, 15, 10, 13, 15])
    ws3 = wb.create_sheet("Summary", 0)
    _title(ws3, "RENTAL ARREARS AGEING REPORT", "As at 30 September 2026 - Finance & Arrears", 4)
    xl_header(ws3, 4, ["Bucket (days)", "No. of tenancies", "Amount (incl GST)", "% of total"])
    for i, b in enumerate(["0-30", "31-60", "61-90", ">90"], 5):
        ws3.cell(i, 1, b)
        ws3.cell(i, 2, f'=COUNTIF(Data!L2:L{n},"{b}")')
        ws3.cell(i, 3, f'=SUMIF(Data!L2:L{n},"{b}",Data!J2:J{n})').number_format = MONEY
        ws3.cell(i, 4, f"=C{i}/C9").number_format = "0.0%"
    ws3.cell(9, 1, "Total").font = Font(bold=True)
    ws3.cell(9, 2, "=SUM(B5:B8)")
    ws3.cell(9, 3, "=SUM(C5:C8)").number_format = MONEY
    xl_header(ws3, 11, ["Town", "No. of tenancies", "Amount (incl GST)", "> 3 months' rent"])
    for i, t in enumerate(TOWNS, 12):
        ws3.cell(i, 1, t["name"])
        ws3.cell(i, 2, f'=COUNTIF(Data!E2:E{n},A{i})')
        ws3.cell(i, 3, f'=SUMIF(Data!E2:E{n},A{i},Data!J2:J{n})').number_format = MONEY
        ws3.cell(i, 4, f'=COUNTIFS(Data!E2:E{n},A{i},Data!M2:M{n},"YES")')
    ws3.cell(19, 1, "Notes: Termination threshold = arrears > 3 months' rent incl. GST (PLG-POL-003 v2.0). Interest not included.").font = Font(italic=True)
    ws3.cell(20, 1, f"Static check: {len(arr)} tenancies in arrears, total S${sum(r['amount'] for r in arr):,.2f}.").font = Font(italic=True)
    xl_widths(ws3, [18, 16, 18, 16])
    xl_props(wb, "Arrears Ageing Sep 2026", sname("S07"), D(2026, 10, 2))
    wb.save(path)


def x_rentroll(M, path):
    asof = D(2026, 9, 30)
    wb = Workbook()
    wu = wb.active
    wu.title = "Units"
    xl_header(wu, 1, ["Unit ID", "Town", "Address", "Unit type", "Floor area (sqm)", "NC", "Status"])
    for i, u in enumerate(M["units"], 2):
        _rows(wu, i, [[u["unit_id"], town_name(u["town"]), addr(u), u["unit_type"], u["floor_area_sqm"], u["nc_id"], u["status"]]])
    nu = len(M["units"]) + 1
    xl_widths(wu, [16, 14, 52, 24, 12, 10, 12])
    ws = wb.create_sheet("Rent Roll", 0)
    hdr = ["Tenancy ID", "Unit ID", "Tenant", "Trade", "Start", "End", "Monthly rent", "Town (lookup)", "Area sqm (lookup)", "Rent psm", "Annual rent", "Deposit held"]
    xl_header(ws, 1, hdr)
    act = sorted([t for t in M["tenancies"] if t["status"] == "Active"], key=lambda t: t["tenancy_id"])
    for i, ta in enumerate(act, 2):
        _rows(ws, i, [[ta["tenancy_id"], ta["unit_id"], tenant(M, ta["tenant_id"])["name"], ta["trade"], ta["start_date"], ta["end_date"], ta["monthly_rent"],
                       f"=VLOOKUP(B{i},Units!$A$2:$G${nu},2,FALSE)", f"=VLOOKUP(B{i},Units!$A$2:$G${nu},5,FALSE)", f"=G{i}/I{i}", f"=G{i}*12", ta["security_deposit"]]],
              {4: DATEF, 5: DATEF, 6: MONEY, 9: MONEY, 10: MONEY, 11: MONEY})
    n = len(act) + 1
    ws.freeze_panes = "A2"
    xl_widths(ws, [15, 16, 34, 26, 12, 12, 13, 14, 12, 10, 14, 14])
    wt = wb.create_sheet("By Town", 1)
    _title(wt, "RENT ROLL FY2026 - SUMMARY BY TOWN", f"Active tenancies as at {asof:%d %b %Y}", 4)
    xl_header(wt, 4, ["Town", "Active tenancies", "Monthly rent", "Annual rent"])
    for i, t in enumerate(TOWNS, 5):
        wt.cell(i, 1, t["name"])
        wt.cell(i, 2, f"=COUNTIF('Rent Roll'!H2:H{n},A{i})")
        wt.cell(i, 3, f"=SUMIF('Rent Roll'!H2:H{n},A{i},'Rent Roll'!G2:G{n})").number_format = MONEY
        wt.cell(i, 4, f"=C{i}*12").number_format = MONEY
    wt.cell(11, 1, "Total").font = Font(bold=True)
    wt.cell(11, 2, "=SUM(B5:B10)")
    wt.cell(11, 3, "=SUM(C5:C10)").number_format = MONEY
    wt.cell(11, 4, "=SUM(D5:D10)").number_format = MONEY
    xl_widths(wt, [18, 16, 16, 18])
    xl_props(wb, "Rent Roll FY2026", sname("S07"), D(2026, 10, 1))
    wb.save(path)


STAFF_ALIASES = {"S06": ["Daniel Ong", "daniel ong", "Daniel O.", "D. Ong"], "S04": ["Marcus Lim", "marcus", "M. Lim"], "S12": ["Benjamin Ho", "Ben Ho", "B. Ho"]}


def x_inspections(M, path):
    rng = random.Random(M["seed"] + 21)
    rows = [i for i in M["inspections"] if i["inspection_date"].year == 2026]
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    hdr = ["Inspection ID", "Date", "Unit", "Town", "Officer", "Type", "Finding", "Outcome", "F/U due", "Closed", "Days to close", "Status"]
    xl_header(ws, 1, hdr)
    for j, i in enumerate(rows, 2):
        off = rng.choice(STAFF_ALIASES.get(i["officer_id"], [sname(i["officer_id"])]))
        _rows(ws, j, [[i["inspection_id"], i["inspection_date"], i["unit_id"], town_name(unit(M, i["unit_id"])["town"]), off, i["inspection_type"], i["finding_code"],
                       i["outcome"], i["follow_up_due"], i["closed_date"], f'=IF(J{j}="","",J{j}-B{j})',
                       f'=IF(H{j}="Compliant","N/A",IF(J{j}<>"","Closed",IF(I{j}<DATE(2026,9,30),"OVERDUE","Open")))']], {1: DATEF, 8: DATEF, 9: DATEF})
    n = len(rows) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{n}"
    xl_widths(ws, [15, 12, 16, 14, 14, 16, 18, 13, 12, 12, 12, 10])
    s = wb.create_sheet("Summary", 0)
    _title(s, "INSPECTION TRACKER 2026", "Counts by month and outcome (formulas on Tracker sheet)", 5)
    xl_header(s, 4, ["Month", "Compliant", "Minor Breach", "Major Breach", "Total"])
    for m in range(1, 10):
        r = 4 + m
        s.cell(r, 1, D(2026, m, 1)).number_format = "mmm-yyyy"
        for c, o in enumerate(["Compliant", "Minor Breach", "Major Breach"], 2):
            s.cell(r, c, f'=COUNTIFS(Tracker!$B$2:$B${n},">="&$A{r},Tracker!$B$2:$B${n},"<"&EDATE($A{r},1),Tracker!$H$2:$H${n},"{o}")')
        s.cell(r, 5, f"=SUM(B{r}:D{r})")
    s.cell(15, 1, "NB officer names typed inconsistently - use Officer ID mapping in PLG Records for reporting.").font = Font(italic=True, color="C00000")
    xl_widths(s, [12, 12, 14, 14, 10])
    xl_props(wb, "Inspection Tracker 2026", sname("S06"), D(2026, 9, 30))
    wb.save(path)


def x_feedback_q3(M, path):
    a, b = D(2026, 7, 1), D(2026, 9, 30)
    fb = [f for f in M["feedback"] if a <= f["received_date"] <= b]
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback Q3"
    xl_header(ws, 1, ["Feedback ID", "Received", "Channel", "Unit", "Category code", "Category", "Severity", "Team", "i-EMS ticket", "Status", "Description"])
    for j, f in enumerate(fb, 2):
        _rows(ws, j, [[f["feedback_id"], f["received_date"], f["channel"], f["unit_id"], f["category_code"], f["category"], f["severity"], f["assigned_team"],
                       f["iems_ticket_id"], f["status"], f["description"]]], {1: DATEF})
    n = len(fb) + 1
    xl_widths(ws, [14, 12, 12, 16, 8, 26, 9, 20, 15, 11, 80])
    s = wb.create_sheet("By Category", 0)
    _title(s, "PUBLIC FEEDBACK & DEFECT STATISTICS - Q3 2026", "Jul-Sep 2026", 5)
    xl_header(s, 4, ["Category", "Jul", "Aug", "Sep", "Q3 total"])
    for r, t in enumerate(DEFECT_TAXONOMY, 5):
        s.cell(r, 1, t["category"])
        for c, m in enumerate([7, 8, 9], 2):
            s.cell(r, c, f"=COUNTIFS('Feedback Q3'!$F$2:$F${n},$A{r},'Feedback Q3'!$B$2:$B${n},\">=\"&DATE(2026,{m},1),'Feedback Q3'!$B$2:$B${n},\"<\"&DATE(2026,{m + 1},1))")
        s.cell(r, 5, f"=SUM(B{r}:D{r})")
    tr = 5 + len(DEFECT_TAXONOMY)
    s.cell(tr, 1, "Total").font = Font(bold=True)
    for c in "BCDE":
        s[f"{c}{tr}"] = f"=SUM({c}5:{c}{tr - 1})"
    s.cell(tr + 1, 1, f"(Static count: {len(fb)} cases received in Q3)").font = Font(italic=True)
    xl_widths(s, [30, 8, 8, 8, 10])
    s2 = wb.create_sheet("i-EMS SLA")
    tk = [t for t in M["tickets"] if a <= t["created_date"] <= b]
    xl_header(s2, 1, ["Team", "Tickets raised", "Closed", "SLA met", "SLA %", "Misrouted"])
    teams = sorted({t["team"] for t in tk})
    for r, team in enumerate(teams, 2):
        tt = [t for t in tk if t["team"] == team]
        cl = [t for t in tt if t["sla_met"] is not None]
        _rows(s2, r, [[team, len(tt), len(cl), sum(1 for t in cl if t["sla_met"]), f"=IF(C{r}=0,\"\",D{r}/C{r})", sum(1 for t in tt if t["status"].startswith("Rejected"))]])
        s2.cell(r, 5).number_format = "0.0%"
    xl_widths(s2, [22, 14, 10, 10, 10, 10])
    xl_props(wb, "Defect & Feedback Stats Q3 2026", sname("S11"), D(2026, 10, 3))
    wb.save(path)


def x_tasks(M, path):
    rng = random.Random(M["seed"] + 31)
    open_t = [t for t in M["tasks"] if t["status"] != "Done"]
    wb = Workbook()
    s = wb.active
    s.title = "Overview"
    s.merge_cells("A1:F1")
    s["A1"] = "PLG OFFICER TASK TRACKER (shared - pls update by Fri!!)"
    s["A1"].font = TITLE_FONT
    s["A2"] = "last updated 28/9 by Nurul. Ahmad's tab still has Aug items??"
    s["A2"].font = Font(italic=True, color="C00000")
    xl_header(s, 4, ["Officer", "Sheet", "Open items", "Overdue (per tracker)"])
    by = defaultdict(list)
    for t in open_t:
        by[t["assignee_id"]].append(t)
    r = 5
    for sid in sorted(by):
        nm = sname(sid)
        sheet = nm.split()[0][:10]
        ws = wb.create_sheet(sheet)
        ws.merge_cells("A1:G1")
        ws["A1"] = f"{nm} - open tasks"
        ws["A1"].font = Font(bold=True, size=12)
        xl_header(ws, 2, ["Task ID", "Type", "Description", "Related ref", "Due", "Status", "Remarks"])
        for j, t in enumerate(sorted(by[sid], key=lambda t: t["due_date"]), 3):
            due = t["due_date"] if rng.random() > 0.12 else t["due_date"].strftime("%d/%m/%y")  # messy: some dates typed as text
            rem = rng.choice(["", "", "", "called tenant, no answer", "waiting for tenant docs", "chase vendor", "pls see email", "escalated to EM", "KIV"])
            _rows(ws, j, [[t["task_id"], t["task_type"], t["title"], t["related_id"], due, t["status"], rem]], {4: DATEF})
        xl_widths(ws, [14, 20, 60, 16, 12, 12, 26])
        s.cell(r, 1, nm)
        s.cell(r, 2, sheet)
        s.cell(r, 3, f"=COUNTA('{sheet}'!A3:A2000)")
        s.cell(r, 4, f"=COUNTIF('{sheet}'!F3:F2000,\"Overdue\")")
        r += 1
    xl_widths(s, [22, 12, 12, 20])
    xl_props(wb, "Officer Task Tracker", sname("S03"), D(2026, 9, 28))
    wb.save(path)


def x_trade_mix(M, path):
    act = [t for t in M["tenancies"] if t["status"] == "Active"]
    wb = Workbook()
    d = wb.active
    d.title = "Active Tenancies"
    xl_header(d, 1, ["Tenancy ID", "Unit ID", "NC", "Trade", "Trade group"])
    for i, ta in enumerate(act, 2):
        _rows(d, i, [[ta["tenancy_id"], ta["unit_id"], unit(M, ta["unit_id"])["nc_id"], ta["trade"], TRADES[ta["trade"]][2]]])
    n = len(act) + 1
    xl_widths(d, [15, 16, 10, 28, 20])
    trades = sorted(TRADES)
    m = wb.create_sheet("Matrix", 0)
    _title(m, "TRADE MIX ANALYSIS BY NEIGHBOURHOOD CENTRE", "Active tenancies as at 30 Sep 2026; limit = 2 per trade per NC (POL-001 cl. 5.1)", len(trades) + 3)
    xl_header(m, 4, ["NC ID", "NC name"] + trades + ["Max same trade", "Flag"])
    for r, nc in enumerate(M["ncs"], 5):
        m.cell(r, 1, nc["nc_id"])
        m.cell(r, 2, nc["name"])
        for c, tr in enumerate(trades, 3):
            m.cell(r, c, f"=COUNTIFS('Active Tenancies'!$C$2:$C${n},$A{r},'Active Tenancies'!$D$2:$D${n},{get_column_letter(c)}$4)")
        lc = get_column_letter(len(trades) + 2)
        m.cell(r, len(trades) + 3, f"=MAX(C{r}:{lc}{r})")
        m.cell(r, len(trades) + 4, f'=IF({get_column_letter(len(trades) + 3)}{r}>2,"REVIEW","")')
    m.freeze_panes = "C5"
    xl_widths(m, [10, 26] + [9] * len(trades) + [10, 8])
    xl_props(wb, "Trade Mix Analysis by NC", sname("S02"), D(2026, 9, 30))
    wb.save(path)


def x_renewals(M, path):
    a, b = D(2026, 10, 1), D(2027, 9, 30)
    exp = sorted([t for t in M["tenancies"] if t["status"] == "Active" and a <= t["end_date"] <= b], key=lambda t: t["end_date"])
    apps = {}
    for ap in M["applications"]:
        if ap["app_type"] == "RENEWAL" and ap["tenancy_id"]:
            apps[ap["tenancy_id"]] = ap
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    _title(ws, "RENEWALS PIPELINE - TENANCIES EXPIRING OCT 2026 TO SEP 2027", "Application deadline = expiry - 6 months (PLG-POL-002 cl. 4.2)", 11)
    xl_header(ws, 4, ["Tenancy ID", "Tenant", "Unit", "Trade", "Expiry", "Monthly rent", "Apply-by", "Application", "Submitted", "Late?", "App status"])
    for i, t in enumerate(exp, 5):
        ap = apps.get(t["tenancy_id"])
        _rows(ws, i, [[t["tenancy_id"], tenant(M, t["tenant_id"])["name"], t["unit_id"], t["trade"], t["end_date"], t["monthly_rent"], f"=EDATE(E{i},-6)",
                       ap["application_id"] if ap else None, ap["submitted_date"] if ap else None, f'=IF(I{i}="","",IF(I{i}>G{i},"LATE",""))',
                       ap["status"] if ap else "No application"]], {4: DATEF, 5: MONEY, 6: DATEF, 8: DATEF})
    ws.freeze_panes = "A5"
    xl_widths(ws, [15, 34, 16, 24, 12, 12, 12, 14, 12, 8, 16])
    xl_props(wb, "Renewals Pipeline FY2026", sname("S05"), D(2026, 9, 25))
    wb.save(path)


# ------------------------------------------------------------------ CSV / JSON exports
def _ser(o):
    return {k: iso(v) for k, v in o.items()}


def j_dump(path, obj):
    with open(path, "w") as f:
        json.dump(obj, f, indent=2, default=iso)


def csv_dump(path, rows, fields=None):
    fields = fields or list(rows[0].keys())
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: iso(r.get(k)) for k in fields})


def e_crm(M, path):
    recs = [f for f in M["feedback"] if D(2026, 7, 1) <= f["received_date"] <= D(2026, 9, 30)]
    j_dump(path, dict(export=dict(system="CRM (Customer Relations)", object="feedback_case", generated_at="2026-10-01T06:00:00+08:00", record_count=len(recs),
                                  filter="received_date between 2026-07-01 and 2026-09-30"),
                      records=[dict(id=f["feedback_id"], received=f["received_date"], channel=f["channel"], location=dict(unit_id=f["unit_id"], block_id=f["block_id"]),
                                    text=f["description"], classification=dict(code=f["category_code"], category=f["category"], severity=f["severity"],
                                                                               team=f["assigned_team"], officer_confidence=f["label_confidence"]),
                                    status=f["status"], iems_ref=f["iems_ticket_id"], officer=f["officer_id"], reply_due=f["reply_due"], replied=f["replied_date"])
                               for f in recs]))


def e_iems(M, path):
    j_dump(path, dict(source="i-EMS", api_version="v2", generated_at="2026-10-01T06:05:00+08:00",
                      tickets=[dict(ticketId=t["ticket_id"], sourceRef=t["feedback_id"], createdOn=t["created_date"], location=t["unit_id"],
                                    defectCode=t["category_code"], defectCategory=TAX_BY_CODE[t["category_code"]]["category"], assignedTeam=t["team"],
                                    contractor=t["vendor"] or None, priority=t["priority"], slaDays=t["sla_days"], status=t["status"], completedOn=t["completed_date"],
                                    slaMet=t["sla_met"], raisedBy=t["raised_by"]) for t in M["tickets"]]))


def e_invoices(M, path):
    rows = [dict(i, tenant_name=tenant(M, i["tenant_id"])["name"]) for i in M["invoices"] if i["period"] == "2026-09"]
    rows += [dict(rows[17]), dict(rows[402])]  # duplicated rows from ERP re-run (data quality edge case)
    csv_dump(path, rows, ["invoice_id", "tenancy_id", "tenant_id", "tenant_name", "period", "invoice_date", "due_date", "rent_amount", "gst_amount", "total_amount",
                          "status", "paid_date"])


def e_checklists(M, path):
    j_dump(path, dict(version="2026.1", owner="Tenancy Management", effective="2026-01-01",
                      checklists={k: dict(label=APP_TYPE_LABEL[k], documents=[dict(code=c, description=d, mandatory=m) for c, d, m in v]) for k, v in CHECKLISTS.items()},
                      rules=["Business profile must be dated within 3 months of submission.", "Insurance certificate must cover the whole tenancy term (or the A&A works period).",
                             "Documents must name the correct unit address/unit ID.", "Illegible scans are treated as not submitted."]))


def e_taxonomy(M, path):
    j_dump(path, dict(version="SOP-020 v2.0", categories=DEFECT_TAXONOMY, severity_sla_days=SEVERITY_SLA_DAYS,
                      guidance="Choose one primary category; if multiple issues, choose the highest severity; flag ambiguous cases for officer review."))


def e_tasks(M, path):
    rows = [t for t in M["tasks"] if t["status"] != "Done"]
    rows = [dict(t, assignee_name=sname(t["assignee_id"])) for t in rows]
    csv_dump(path, rows, ["task_id", "task_type", "title", "related_id", "assignee_id", "assignee_name", "created_date", "due_date", "status", "safety_flag",
                          "revenue_at_risk", "tenants_affected"])


def e_contacts(M, path):
    rng = random.Random(M["seed"] + 41)
    rows = []
    active_tenants = {t["tenant_id"] for t in M["tenancies"] if t["status"] in ("Active", "Pending Commencement")}
    for t in M["tenants"]:
        if t["tenant_id"] not in active_tenants:
            continue
        nm = t["name"]
        if rng.random() < 0.08:
            nm = nm.upper()
        if rng.random() < 0.05:
            nm = nm.replace("Pte Ltd", "Pte. Ltd.")
        rows.append(dict(crm_contact_id=f"C{int(t['tenant_id'][2:]) * 7 % 100000:05d}", tenant_id=t["tenant_id"], business_name=nm, contact_person=t["contact_person"],
                         email=t["email"] if rng.random() > 0.04 else "", phone=t["phone"], preferred_channel=rng.choice(["Email", "Email", "SMS", "Post"])))
    kh = next(r for r in rows if r["business_name"].lower().startswith("kim hiap"))
    rows.append(dict(kh, crm_contact_id="C99871", tenant_id="", business_name="Kim Hiap Trading Pte. Ltd.", email="kimhiap.trading@example.com"))
    csv_dump(path, rows)


# ------------------------------------------------------------------ database export
SCHEMA = """-- PLG Records (fictional) - relational schema for the POC database export
-- Engine: PostgreSQL-compatible DDL. CSV files in this folder are one per table.

CREATE TABLE staff (staff_id VARCHAR(4) PRIMARY KEY, name TEXT NOT NULL, title TEXT, dept TEXT, email TEXT UNIQUE, phone TEXT);
CREATE TABLE towns (town_code CHAR(2) PRIMARY KEY, town_name TEXT NOT NULL, region TEXT, town_council TEXT, estate_manager_id VARCHAR(4) REFERENCES staff);
CREATE TABLE neighbourhood_centres (nc_id VARCHAR(8) PRIMARY KEY, town_code CHAR(2) REFERENCES towns, nc_name TEXT);
CREATE TABLE blocks (block_id VARCHAR(8) PRIMARY KEY, town_code CHAR(2) REFERENCES towns, blk INT, street TEXT, nc_id VARCHAR(8) REFERENCES neighbourhood_centres,
                     postal_code CHAR(6), year_built INT);
CREATE TABLE trades (trade TEXT PRIMARY KEY, trade_group TEXT, allowed_unit_types TEXT);
CREATE TABLE units (unit_id VARCHAR(16) PRIMARY KEY, block_id VARCHAR(8) REFERENCES blocks, unit_no VARCHAR(8), unit_type TEXT, floor_area_sqm NUMERIC(6,1),
                    status TEXT CHECK (status IN ('Tenanted','Vacant','Under Tender','Not Lettable')), vacant_since DATE);
CREATE TABLE tenants (tenant_id VARCHAR(8) PRIMARY KEY, name TEXT NOT NULL, brn VARCHAR(12), entity_type TEXT, contact_person TEXT, email TEXT, phone TEXT,
                      registered_since DATE);
CREATE TABLE tenancies (tenancy_id VARCHAR(16) PRIMARY KEY, unit_id VARCHAR(16) REFERENCES units, tenant_id VARCHAR(8) REFERENCES tenants,
                        trade TEXT REFERENCES trades, start_date DATE, end_date DATE, monthly_rent NUMERIC(10,2), security_deposit NUMERIC(10,2),
                        status TEXT, rent_basis TEXT);
CREATE TABLE rent_invoices (invoice_id VARCHAR(16) PRIMARY KEY, tenancy_id VARCHAR(16) REFERENCES tenancies, period CHAR(7), invoice_date DATE, due_date DATE,
                            rent_amount NUMERIC(10,2), gst_amount NUMERIC(10,2), total_amount NUMERIC(10,2), status TEXT);
CREATE TABLE payments (payment_id VARCHAR(12) PRIMARY KEY, invoice_id VARCHAR(16) REFERENCES rent_invoices, tenancy_id VARCHAR(16) REFERENCES tenancies,
                       paid_date DATE, amount NUMERIC(10,2), method TEXT);
CREATE TABLE inspections (inspection_id VARCHAR(12) PRIMARY KEY, unit_id VARCHAR(16) REFERENCES units, tenancy_id VARCHAR(16) REFERENCES tenancies,
                          inspection_date DATE, officer_id VARCHAR(4) REFERENCES staff, inspection_type TEXT, finding_code TEXT, outcome TEXT,
                          follow_up_due DATE, closed_date DATE, remarks TEXT);
CREATE TABLE feedback_cases (feedback_id VARCHAR(14) PRIMARY KEY, received_date DATE, channel TEXT, unit_id VARCHAR(16) REFERENCES units, description TEXT,
                             category_code CHAR(3), severity TEXT, assigned_team TEXT, status TEXT, iems_ticket_id VARCHAR(16), officer_id VARCHAR(4) REFERENCES staff,
                             reply_due DATE, replied_date DATE);
CREATE TABLE iems_tickets (ticket_id VARCHAR(16) PRIMARY KEY, feedback_id VARCHAR(14) REFERENCES feedback_cases, created_date DATE, unit_id VARCHAR(16) REFERENCES units,
                           category_code CHAR(3), team TEXT, vendor TEXT, priority TEXT, sla_days INT, status TEXT, completed_date DATE, sla_met BOOLEAN);
CREATE TABLE applications (application_id VARCHAR(12) PRIMARY KEY, app_type TEXT, tenancy_id VARCHAR(16) REFERENCES tenancies, tenant_id VARCHAR(8) REFERENCES tenants,
                           unit_id VARCHAR(16) REFERENCES units, submitted_date DATE, status TEXT, officer_id VARCHAR(4) REFERENCES staff, decision_date DATE,
                           proposed_trade TEXT, remarks TEXT);
CREATE TABLE application_documents (application_id VARCHAR(12) REFERENCES applications, doc_code VARCHAR(8), doc_label TEXT, mandatory BOOLEAN, submitted BOOLEAN,
                                    received_date DATE, verified BOOLEAN, PRIMARY KEY (application_id, doc_code));
CREATE TABLE tasks (task_id VARCHAR(12) PRIMARY KEY, task_type TEXT, title TEXT, related_id VARCHAR(16), assignee_id VARCHAR(4) REFERENCES staff, created_date DATE,
                    due_date DATE, status TEXT, safety_flag BOOLEAN, revenue_at_risk NUMERIC(12,2), tenants_affected INT);
CREATE TABLE letters (letter_id VARCHAR(14) PRIMARY KEY, letter_date DATE, template_id VARCHAR(12), letter_type TEXT, reference_id VARCHAR(16),
                      tenant_id VARCHAR(8) REFERENCES tenants, author_id VARCHAR(4) REFERENCES staff, channel TEXT);

CREATE INDEX ix_units_block ON units(block_id);
CREATE INDEX ix_units_status ON units(status);
CREATE INDEX ix_tenancies_unit ON tenancies(unit_id);
CREATE INDEX ix_tenancies_tenant ON tenancies(tenant_id);
CREATE INDEX ix_tenancies_end ON tenancies(end_date);
CREATE INDEX ix_invoices_tenancy ON rent_invoices(tenancy_id, due_date);
CREATE INDEX ix_payments_invoice ON payments(invoice_id);
CREATE INDEX ix_inspections_unit ON inspections(unit_id, inspection_date);
CREATE INDEX ix_feedback_unit ON feedback_cases(unit_id, received_date);
CREATE INDEX ix_tickets_feedback ON iems_tickets(feedback_id);
CREATE INDEX ix_apps_tenant ON applications(tenant_id);
CREATE INDEX ix_tasks_assignee ON tasks(assignee_id, status, due_date);

-- NOTE: the export intentionally contains data-quality issues for testing (see _report.md):
--   duplicate tenant (same BRN), NULL emails/floor areas, an outlier rent, a far-future inspection date, a duplicated payment row.
--   The duplicated payment row violates the payments PK on purpose; load with a staging table or DuckDB read_csv.
"""


def db_tables(M):
    rng = random.Random(M["seed"] + 51)
    T = {}
    T["staff"] = [dict(staff_id=s["id"], name=s["name"], title=s["title"], dept=s["dept"], email=s["email"], phone=s["phone"]) for s in STAFF]
    T["towns"] = [dict(town_code=t["code"], town_name=t["name"], region=t["region"], town_council=t["tc"], estate_manager_id=t["em"]) for t in TOWNS]
    T["neighbourhood_centres"] = [dict(nc_id=n["nc_id"], town_code=n["town"], nc_name=n["name"]) for n in M["ncs"]]
    T["blocks"] = [dict(block_id=b["block_id"], town_code=b["town"], blk=b["blk"], street=b["street"], nc_id=b["nc_id"], postal_code=b["postal"], year_built=b["year_built"])
                   for b in M["blocks"]]
    T["trades"] = [dict(trade=k, trade_group=v[2], allowed_unit_types="|".join(v[1])) for k, v in TRADES.items()]
    units = [dict(unit_id=u["unit_id"], block_id=u["block_id"], unit_no=u["unit_no"], unit_type=u["unit_type"], floor_area_sqm=u["floor_area_sqm"], status=u["status"],
                  vacant_since=u.get("vacant_since")) for u in M["units"]]
    for u in rng.sample([u for u in units if u["unit_id"] not in ("U-BV214-01-32",)], 5):
        u["floor_area_sqm"] = None  # NULL edge case
    T["units"] = units
    tenants = [dict(tenant_id=t["tenant_id"], name=t["name"], brn=t["brn"], entity_type=t["entity_type"], contact_person=t["contact_person"],
                    email=t["email"] if rng.random() > 0.03 else None, phone=t["phone"], registered_since=t["registered_since"]) for t in M["tenants"]]
    kh = next(t for t in tenants if t["tenant_id"] == M["key"]["Kim Hiap Trading"]["tenant"])
    tenants.append(dict(kh, tenant_id="T-19807", name="Kim Hiap Trading Pte. Ltd.", email="kimhiap.trading@example.com"))  # duplicate entity, same BRN
    T["tenants"] = tenants
    tens = [dict(tenancy_id=t["tenancy_id"], unit_id=t["unit_id"], tenant_id=t["tenant_id"], trade=t["trade"], start_date=t["start_date"], end_date=t["end_date"],
                 monthly_rent=t["monthly_rent"], security_deposit=t["security_deposit"], status=t["status"], rent_basis=t["rent_basis"]) for t in M["tenancies"]]
    hist = [t for t in tens if t["status"] in ("Expired", "Terminated")]
    out = rng.choice(hist)
    out["monthly_rent"] = out["monthly_rent"] * 10  # outlier (extra zero keyed in on a historical record)
    T["tenancies"] = tens
    T["rent_invoices"] = [dict(invoice_id=i["invoice_id"], tenancy_id=i["tenancy_id"], period=i["period"], invoice_date=i["invoice_date"], due_date=i["due_date"],
                               rent_amount=i["rent_amount"], gst_amount=i["gst_amount"], total_amount=i["total_amount"], status=i["status"]) for i in M["invoices"]]
    pays = [dict(p) for p in M["payments"]]
    pays.append(dict(pays[1234]))  # duplicated payment row
    T["payments"] = pays
    ins = [dict(i) for i in M["inspections"]]
    bad = rng.choice([i for i in ins if i["finding_code"] == "NONE"])
    bad["inspection_date"] = D(2099, bad["inspection_date"].month, bad["inspection_date"].day)  # impossible date
    T["inspections"] = ins
    T["feedback_cases"] = [{k: f[k] for k in ["feedback_id", "received_date", "channel", "unit_id", "description", "category_code", "severity", "assigned_team",
                                               "status", "iems_ticket_id", "officer_id", "reply_due", "replied_date"]} for f in M["feedback"]]
    T["iems_tickets"] = [{k: t[k] for k in ["ticket_id", "feedback_id", "created_date", "unit_id", "category_code", "team", "vendor", "priority", "sla_days", "status",
                                             "completed_date", "sla_met"]} for t in M["tickets"]]
    T["applications"] = [dict(a) for a in M["applications"]]
    T["application_documents"] = [dict(d) for d in M["app_docs"]]
    T["tasks"] = [dict(t) for t in M["tasks"]]
    T["letters"] = [dict(l) for l in M["letters"]]
    dq = dict(null_floor_area=[u["unit_id"] for u in units if u["floor_area_sqm"] is None], duplicate_tenant=dict(original=kh["tenant_id"], duplicate="T-19807", brn=kh["brn"]),
              outlier_rent=dict(tenancy_id=out["tenancy_id"], monthly_rent=out["monthly_rent"]), impossible_date=dict(inspection_id=bad["inspection_id"], date=bad["inspection_date"]),
              duplicate_payment=pays[-1]["payment_id"], null_tenant_emails=sum(1 for t in tenants if t["email"] is None))
    return T, dq


def register_data(M):
    xl = [
        ("XLS-001", "Vacancy_Report_Aug2026.xlsx", "Monthly Vacancy Report - Aug 2026", "S08", D(2026, 9, 2), x_vacancy,
         "Vacancy by town with merged headers, formulas and notes rows (47 vacant)", "How many vacant units in each town at end-Aug 2026?", "CONFLICT C4 (vs DECK-001); merged headers; notes rows"),
        ("XLS-002", "Arrears_Ageing_Sep2026.xlsx", "Rental Arrears Ageing Report - Sep 2026", "S07", D(2026, 10, 2), x_arrears,
         "Arrears by tenancy with ageing buckets, SUMIF/COUNTIF summary", "Which tenants owe more than 3 months' rent?; Total arrears in the >90 bucket?", ""),
        ("XLS-003", "Rent_Roll_FY2026.xlsx", "Rent Roll FY2026", "S07", D(2026, 10, 1), x_rentroll,
         "Active tenancies with VLOOKUP to units, annual rent and deposit", "What is total monthly rent by town?; What rent is recorded for Heng Ann Confectionery?", "CONFLICT C6 (Heng Ann rent 4,850)"),
        ("XLS-004", "Inspection_Tracker_2026.xlsx", "Inspection Tracker 2026", "S06", D(2026, 9, 30), x_inspections,
         "2026 inspections; inconsistent officer names; COUNTIFS summary", "How many major breaches were found in 2026?; Which follow-ups are overdue?", "Messy: inconsistent officer naming"),
        ("XLS-005", "Defect_Feedback_Stats_Q3_2026.xlsx", "Defect & Feedback Statistics Q3 2026", "S11", D(2026, 10, 3), x_feedback_q3,
         "Q3 feedback by category/month and i-EMS SLA by team", "What was the most common feedback category in Q3 2026?; SLA % for M&E?", ""),
        ("XLS-006", "Officer_Task_Tracker.xlsx", "PLG Officer Task Tracker", "S03", D(2026, 9, 28), x_tasks,
         "Shared messy tracker: per-officer sheets, dates typed as text, remarks", "What open tasks does Siti have?", "Messy: text dates, notes rows, merged title"),
        ("XLS-007", "Trade_Mix_Analysis_by_NC.xlsx", "Trade Mix Analysis by NC", "S02", D(2026, 9, 30), x_trade_mix,
         "COUNTIFS matrix NC x trade with limit flag", f"How many bubble tea outlets are in {M['key']['bubble_tea_nc']['nc_id']}?", ""),
        ("XLS-008", "Renewals_Pipeline_FY2026.xlsx", "Renewals Pipeline Oct 2026 - Sep 2027", "S05", D(2026, 9, 25), x_renewals,
         "Expiring tenancies with apply-by date (EDATE) and LATE flag", "Which renewals were submitted late?; How many tenancies expire in Q4 2026?", "Relates to C3"),
    ]
    for id, fname, title, author, date, fn_, purpose, q, flags in xl:
        register(id, f"sheets/{fname}", "xlsx", title, sname(author), date, purpose, q, (lambda root, fn_=fn_, fname=fname: fn_(M, f"{root}/sheets/{fname}")), flags=flags)
    ex = [
        ("EXP-001", "crm_feedback_export_2026Q3.json", "CRM feedback export Q3 2026", "CRM", D(2026, 10, 1), e_crm, "Labelled public feedback cases (defect classification golden data)",
         "Classify feedback FB-xxxx; how many ambiguous cases?"),
        ("EXP-002", "iems_tickets_export.json", "i-EMS tickets export", "i-EMS", D(2026, 10, 1), e_iems, "All i-EMS tickets with SLA results", "Which tickets missed SLA?; status of the Blk 214 seepage ticket?"),
        ("EXP-003", "erp_rent_invoices_2026-09.csv", "ERP rent invoices Sep 2026", "ERP", D(2026, 10, 1), e_invoices, "September invoices with payment status (2 duplicated rows)",
         "How many September invoices are unpaid?"),
        ("EXP-005", "submission_checklists.json", "Submission checklists by application type", "Tenancy Management", D(2026, 1, 1), e_checklists,
         "Checklist master for document-completeness workflow", "What documents are required for an A&A application?"),
        ("EXP-006", "defect_taxonomy.json", "Defect taxonomy (SOP-020)", "Customer Relations", D(2025, 12, 1), e_taxonomy, "Taxonomy, teams, SLAs for classifier",
         "Which team handles pest cases?"),
        ("EXP-007", "plg_tasks_open.csv", "PLG Records open tasks export", "PLG Records", D(2026, 9, 30), e_tasks, "Open/overdue tasks with priority inputs",
         "What are the top 5 priority tasks?"),
        ("EXP-008", "tenant_contacts_crm.csv", "CRM tenant contacts", "CRM", D(2026, 9, 30), e_contacts, "Tenant contact list with inconsistent names and a duplicate",
         "What is the email for Wellcare Family Clinic?"),
    ]
    for id, fname, title, author, date, fn_, purpose, q in ex:
        flags = "Duplicate rows" if id == "EXP-003" else ("Inconsistent naming; duplicate entity" if id == "EXP-008" else "")
        register(id, f"exports/{fname}", fname.rsplit(".", 1)[1], title, author, date, purpose, q, (lambda root, fn_=fn_, fname=fname: fn_(M, f"{root}/exports/{fname}")),
                 flags=flags)
    # database export
    T, dq = db_tables(M)
    M["data_quality"] = dq

    def schema_fn(root):
        with open(f"{root}/database/schema.sql", "w") as f:
            f.write(SCHEMA)
    register("DB-000", "database/schema.sql", "sql", "PLG Records schema (DDL)", "PLG Records", D(2026, 9, 30), "Schema with FKs and indexes for the DB export",
             "How do tenancies join to units and invoices?", schema_fn)
    for i, (tname, rows) in enumerate(T.items(), 1):
        def fn(root, tname=tname, rows=rows):
            csv_dump(f"{root}/database/{tname}.csv", rows)
        register(f"DB-{i:03d}", f"database/{tname}.csv", "csv", f"PLG Records table: {tname} ({len(rows)} rows)", "PLG Records", D(2026, 9, 30),
                 f"Database table export ({len(rows)} rows)", f"Query {tname} via NL->SQL", fn,
                 flags={"tenants": "Duplicate tenant; NULL emails", "units": "NULL floor areas", "tenancies": "Outlier rent; CONFLICT C6", "payments": "Duplicate row",
                        "inspections": "Impossible date (2099)"}.get(tname, ""))
