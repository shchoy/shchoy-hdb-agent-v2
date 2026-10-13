"""Format renderers. Content is expressed as a list of blocks so the same text can go to DOCX or PDF.

Block forms:
  ("h1"|"h2"|"h3", text)  ("p", text)  ("bullets", [..])  ("numbered", [..])
  ("table", header, rows[, col_widths_cm])  ("kv", [(k, v), ..])  ("pagebreak",)  ("note", text)  ("sig", [lines])
"""
import datetime as dt
import hashlib
import io
import random
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches
from pptx.util import Pt as PPt
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from master import GROUP, ORG, ORG_ADDRESS

rl_config.invariant = 1  # reproducible PDF bytes (no timestamps / random IDs)
FIXED_TS = dt.datetime(2026, 9, 30, 18, 0, 0)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def _dt(d):
    return d if isinstance(d, dt.datetime) else dt.datetime(d.year, d.month, d.day, 9, 0)


# ------------------------------------------------------------------ DOCX
def _field(par, code):
    run = par.add_run()
    for tag, txt in (("begin", None), (None, code), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _docx_table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(9)
        _shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def docx_render(path, blocks, *, title, doc_id, author, date, control=None, revisions=None, header_note="", watermark=None, subtitle=None):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    hp = sec.header.paragraphs[0]
    hp.text = f"{ORG} | {GROUP}\t\t{doc_id}" + (f"  [{watermark}]" if watermark else "")
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    fp = sec.footer.paragraphs[0]
    fp.text = f"RESTRICTED - FICTIONAL POC DATA {header_note}   Page "
    _field(fp, "PAGE")
    fp.add_run(" of ")
    _field(fp, "NUMPAGES")
    for r in fp.runs:
        r.font.size = Pt(8)
    h = doc.add_heading(title, level=0)
    for r in h.runs:
        r.font.color.rgb = NAVY
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].italic = True
    if watermark:
        p = doc.add_paragraph()
        r = p.add_run(f"*** {watermark} ***")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    if control:
        doc.add_heading("Document Control", level=2)
        _docx_table(doc, ["Field", "Value"], control, [5, 11])
    if revisions:
        doc.add_heading("Revision History", level=2)
        _docx_table(doc, ["Version", "Date", "Author", "Summary of changes"], revisions, [2, 2.8, 3.5, 8])
    _docx_blocks(doc, blocks)
    cp = doc.core_properties
    cp.author = author
    cp.title = title
    cp.subject = doc_id
    cp.created = _dt(date)
    cp.modified = _dt(date)
    cp.last_modified_by = author
    cp.revision = 1
    doc.save(path)


def _docx_blocks(doc, blocks):
    for b in blocks:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            doc.add_heading(b[1], level=int(k[1]))
        elif k == "p":
            doc.add_paragraph(b[1])
        elif k == "note":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.italic = True
            r.font.size = Pt(9)
        elif k == "bullets":
            for x in b[1]:
                doc.add_paragraph(x, style="List Bullet")
        elif k == "numbered":
            for x in b[1]:
                doc.add_paragraph(x, style="List Number")
        elif k == "table":
            _docx_table(doc, b[1], b[2], b[3] if len(b) > 3 else None)
        elif k == "kv":
            _docx_table(doc, ["Item", "Detail"], b[1], [5, 11])
        elif k == "pagebreak":
            doc.add_page_break()
        elif k == "sig":
            doc.add_paragraph()
            for line in b[1]:
                doc.add_paragraph(line)
        else:
            raise ValueError(k)


# ------------------------------------------------------------------ PDF
_ss = getSampleStyleSheet()
PS = {
    "p": ParagraphStyle("p", parent=_ss["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13),
    "h1": ParagraphStyle("h1", parent=_ss["Heading1"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#1F3A5F"), spaceAfter=8),
    "h2": ParagraphStyle("h2", parent=_ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#1F3A5F")),
    "h3": ParagraphStyle("h3", parent=_ss["Heading3"], fontName="Helvetica-Bold", fontSize=10.5),
    "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10),
    "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=8, leading=10),
    "note": ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=8.5, leading=11, textColor=colors.HexColor("#444444")),
    "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, leading=22, alignment=TA_CENTER, textColor=colors.HexColor("#1F3A5F"), spaceAfter=10),
}


def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class _NumberedCanvas(rl_canvas.Canvas):
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self.setFont("Helvetica", 7.5)
            self.setFillColor(colors.grey)
            self.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {self._pageNumber} of {n}")
            super().showPage()
        super().save()


def _pdf_table(header, rows, widths=None, total_w=17 * cm):
    data = [[Paragraph(esc(h), PS["cellb"]) for h in header]]
    for r in rows:
        data.append([Paragraph(esc("" if v is None else v), PS["cell"]) for v in r])
    if widths:
        cw = [w * cm for w in widths]
    else:
        cw = [total_w / len(header)] * len(header)
    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    return t


def pdf_blocks(blocks):
    story = []
    for b in blocks:
        k = b[0]
        if k in ("h1", "h2", "h3", "p", "note", "title"):
            story.append(Paragraph(esc(b[1]).replace("\n", "<br/>"), PS[k]))
            if k == "p":
                story.append(Spacer(1, 3))
        elif k in ("bullets", "numbered"):
            items = [ListItem(Paragraph(esc(x), PS["p"]), leftIndent=12) for x in b[1]]
            story.append(ListFlowable(items, bulletType="bullet" if k == "bullets" else "1", start=None if k == "bullets" else 1, leftIndent=14,
                                      bulletFontSize=8))
            story.append(Spacer(1, 4))
        elif k == "table":
            story.append(_pdf_table(b[1], b[2], b[3] if len(b) > 3 else None))
            story.append(Spacer(1, 8))
        elif k == "kv":
            story.append(_pdf_table(["Item", "Detail"], b[1], [5, 12]))
            story.append(Spacer(1, 8))
        elif k == "pagebreak":
            story.append(PageBreak())
        elif k == "sig":
            story.append(Spacer(1, 14))
            story.append(KeepTogether([Paragraph(esc(x), PS["p"]) for x in b[1]]))
        elif k == "spacer":
            story.append(Spacer(1, b[1]))
        else:
            raise ValueError(k)
    return story


def pdf_render(path, blocks, *, title, author, subject="", letterhead=True, ref=None, classification="RESTRICTED - FICTIONAL POC DATA"):
    def on_page(c, d):
        c.saveState()
        if letterhead:
            c.setFillColor(colors.HexColor("#1F3A5F"))
            c.rect(0, A4[1] - 2.3 * cm, A4[0], 2.3 * cm, fill=1, stroke=0)
            c.setFillColor(colors.white)
            c.setFont("Helvetica-Bold", 14)
            c.drawString(2 * cm, A4[1] - 1.15 * cm, ORG.upper())
            c.setFont("Helvetica", 8.5)
            c.drawString(2 * cm, A4[1] - 1.7 * cm, f"{GROUP}  |  {ORG_ADDRESS}  |  www.heb.example.org")
            if ref:
                c.drawRightString(A4[0] - 2 * cm, A4[1] - 1.15 * cm, ref)
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.grey)
        c.drawString(2 * cm, 1.2 * cm, classification)
        c.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=3.0 * cm if letterhead else 2 * cm, bottomMargin=2 * cm,
                            title=title, author=author, subject=subject, creator="HEB PLG document system", producer="HEB PLG")
    doc.build(pdf_blocks(blocks), onFirstPage=on_page, onLaterPages=on_page, canvasmaker=_NumberedCanvas)


_FONT_CANDIDATES = ["/System/Library/Fonts/Supplemental/Courier New.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", "/Library/Fonts/Arial.ttf"]


def _font(size):
    for f in _FONT_CANDIDATES:
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def scan_page_image(lines, seed, heavy=False, stamp=None):
    """Render text lines into a noisy, slightly rotated greyscale 'scan'."""
    rng = random.Random(seed)
    W, H = 1240, 1754
    img = Image.new("L", (W, H), 246)
    d = ImageDraw.Draw(img)
    y = 110
    for ln in lines:
        size = 34 if ln.startswith("#") else 25
        txt = ln.lstrip("#").strip()
        d.text((110, y), txt, fill=rng.randint(25, 60), font=_font(size))
        y += int(size * 1.55)
        if y > H - 120:
            break
    if stamp:
        d.ellipse((820, H - 420, 1100, H - 200), outline=90, width=5)
        d.text((850, H - 330), stamp, fill=80, font=_font(26))
    px = img.load()
    for _ in range(9000 if heavy else 2500):
        px[rng.randrange(W), rng.randrange(H)] = rng.randint(120, 200)
    img = img.rotate(rng.uniform(-1.2, 1.2) * (2.5 if heavy else 1), fillcolor=235, resample=Image.BICUBIC)
    if heavy:
        img = img.filter(ImageFilter.GaussianBlur(2.2))
    else:
        img = img.filter(ImageFilter.GaussianBlur(0.6))
    return img


def pdf_scanned(path, pages, *, title, seed, heavy=False, stamps=None):
    c = rl_canvas.Canvas(path, pagesize=A4)
    c.setTitle(title)
    c.setAuthor("Scanner HEB-MFP-07")
    c.setCreator("HEB-MFP-07 Scan to PDF")
    for i, lines in enumerate(pages):
        img = scan_page_image(lines, seed + i, heavy=heavy, stamp=(stamps or {}).get(i))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=55)
        buf.seek(0)
        c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
        c.showPage()
    c.save()


# ------------------------------------------------------------------ EMAIL
def make_eml(path, *, frm, to, cc=None, date, subject, body, msg_id, in_reply_to=None, references=None, attachments=None, extra_headers=None):
    m = EmailMessage()
    m["From"] = frm
    m["To"] = ", ".join(to)
    if cc:
        m["Cc"] = ", ".join(cc)
    m["Date"] = format_datetime(date)
    m["Subject"] = subject
    m["Message-ID"] = msg_id
    if in_reply_to:
        m["In-Reply-To"] = in_reply_to
        m["References"] = " ".join(references or [in_reply_to])
    m["X-Mailer"] = "Lotus Notes Release 9.0.1FP10 (fictional)"
    for k, v in (extra_headers or {}).items():
        m[k] = v
    m.set_content(body)
    for apath, fname, mime in attachments or []:
        with open(apath, "rb") as f:
            data = f.read()
        maintype, subtype = mime.split("/")
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=fname)
    with open(path, "wb") as f:
        f.write(bytes(m))


def msgid(key, domain="heb.example.org"):
    h = hashlib.sha1(key.encode()).hexdigest()[:20]
    return f"<{h}.{key.split('/')[0][:12]}@{domain}>"


# ------------------------------------------------------------------ PPTX
def deck(path, slides, *, title, author, date):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, s in enumerate(slides):
        if i == 0:
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.shapes.title.text = s["title"]
            sl.placeholders[1].text = s.get("subtitle", "")
        else:
            sl = prs.slides.add_slide(prs.slide_layouts[5])  # title only
            sl.shapes.title.text = s["title"]
            sl.shapes.title.text_frame.paragraphs[0].font.size = PPt(30)
            top = Inches(1.5)
            if s.get("bullets"):
                w = Inches(6.0) if (s.get("table") or s.get("chart")) else Inches(12.0)
                tb = sl.shapes.add_textbox(Inches(0.6), top, w, Inches(5.5))
                tf = tb.text_frame
                tf.word_wrap = True
                for j, bl in enumerate(s["bullets"]):
                    lvl = 1 if bl.startswith("- ") else 0
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.text = ("• " if lvl == 0 else "– ") + bl.lstrip("- ")
                    p.level = lvl
                    p.font.size = PPt(18 if lvl == 0 else 15)
            left = Inches(6.9) if s.get("bullets") else Inches(0.6)
            width = Inches(5.9) if s.get("bullets") else Inches(12.0)
            if s.get("table"):
                hdr, rows = s["table"]
                shp = sl.shapes.add_table(len(rows) + 1, len(hdr), left, top, width, Inches(0.35) * (len(rows) + 1))
                tbl = shp.table
                for c, h in enumerate(hdr):
                    tbl.cell(0, c).text = str(h)
                for r, row in enumerate(rows, 1):
                    for c, v in enumerate(row):
                        tbl.cell(r, c).text = str(v)
                for r in range(len(rows) + 1):
                    for c in range(len(hdr)):
                        for p in tbl.cell(r, c).text_frame.paragraphs:
                            p.font.size = PPt(12)
            if s.get("chart"):
                ch = s["chart"]
                cd = CategoryChartData()
                cd.categories = ch["categories"]
                for name, vals in ch["series"].items():
                    cd.add_series(name, vals)
                ctype = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
                         "pie": XL_CHART_TYPE.PIE}[ch.get("type", "bar")]
                gf = sl.shapes.add_chart(ctype, left, top, width, Inches(5.2), cd)
                gf.chart.has_legend = len(ch["series"]) > 1 or ch.get("type") == "pie"
                if gf.chart.has_legend:
                    gf.chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    gf.chart.legend.include_in_layout = False
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    cp = prs.core_properties
    cp.author = author
    cp.title = title
    cp.created = _dt(date)
    cp.modified = _dt(date)
    cp.last_modified_by = author
    prs.save(path)


# ------------------------------------------------------------------ XLSX helpers
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
TITLE_FONT = Font(bold=True, size=14, color="1F3A5F")
THIN = Side(style="thin", color="999999")
MONEY = '"S$"#,##0.00'
DATEF = "dd-mmm-yyyy"


def xl_header(ws, row, headers, col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.border = Border(bottom=THIN, top=THIN, left=THIN, right=THIN)
        c.alignment = Alignment(wrap_text=True, vertical="center")


def xl_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def xl_props(wb, title, author, date):
    wb.properties.title = title
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.properties.created = _dt(date)
    wb.properties.modified = _dt(date)
