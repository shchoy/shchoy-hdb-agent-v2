"""Approved letter templates: the `TPL-*.docx` files of the corpus (SharePoint). Read-only.

A template file has a document-control block (title, metadata tables) followed by the letter, whose slots are
`{{ name }}`. The approved wording is never changed; text in [square brackets] is guidance for officers.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

SLOT = re.compile(r"\{\{\s*(\w+)\s*\}\}")
GUIDANCE = re.compile(r"^\s*\[.*\]\s*$", re.DOTALL)
APPROVED = "Approved for use"


@dataclass(frozen=True)
class Template:
    template_id: str
    title: str
    version: str
    purpose: str
    slots: list[str]
    body: list[str]
    """The letter's paragraphs, guidance removed."""
    path: Path

    def wording(self) -> str:
        return "\n\n".join(self.body)


def letter_paragraphs(document: DocxDocument) -> list[Paragraph]:
    """The letter's own paragraphs: after the document-control block, without officer guidance."""
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Template metadata") + 1
    letter = [p for p in paragraphs[start:] if not GUIDANCE.match(p.text)]
    while letter and not letter[0].text.strip():
        letter.pop(0)
    return letter


def load_template(path: Path) -> Template | None:
    """The template in `path`, or None if it isn't an approved template."""
    document = docx.Document(str(path))
    control = {row.cells[0].text: row.cells[1].text for table in document.tables for row in table.rows[1:]}
    if control.get("Status") != APPROVED:
        return None
    body = [p.text for p in letter_paragraphs(document)]
    slots = [s.strip() for s in control["Required slots"].split(",")]
    return Template(
        template_id=control["Template ID"], title=document.paragraphs[0].text.removesuffix(" (Approved Template)").removeprefix(control["Template ID"]).strip(),
        version=control["Version"], purpose=control["Purpose"], slots=slots, body=body, path=path,
    )


class TemplateLibrary:
    def __init__(self, corpus_dir: Path) -> None:
        self._docs = corpus_dir / "docs"

    def all(self) -> list[Template]:
        found = (load_template(p) for p in sorted(self._docs.glob("TPL-*.docx")))
        return [t for t in found if t is not None]

    def get(self, template_id: str) -> Template | None:
        return next((t for t in self.all() if t.template_id == template_id), None)
