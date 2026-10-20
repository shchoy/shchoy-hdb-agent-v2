"""Drafting through plg-api: the router sends "draft me a letter" to the drafting agent, which picks an approved
template and fills its slots from PLG Records. The officer gets an HTML preview (a `draft` event), a DOCX download,
the slots that couldn't be filled and any personal data found in the draft. Nothing is sent.

Claude is scripted; the records it "looks up" come from the real poc-data/plg_records.db.
"""

from __future__ import annotations

import io
import sqlite3
from typing import Any

import docx

from plg.config import Settings
from plg.llm import ScriptedLLMClient, text, tool_call, turn
from tests.conftest import USER, Api, route

RECORDS = Settings().records_db
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

REMINDER = "Draft a first reminder letter for Golden Leaf Minimart Pte Ltd's overdue rent"
NOTICE = "Draft a missing documents notice for APP-2025-0012"
REPLY = "Draft an update reply to the member of public who gave feedback FB-2025-00002"

INVOICES = "SELECT invoice_id, period, due_date, total_amount FROM rent_invoices WHERE tenancy_id = 'TA-2024-00542' AND status = 'Unpaid' ORDER BY due_date"


def records(sql: str) -> list[tuple[Any, ...]]:
    with sqlite3.connect(f"file:{RECORDS}?mode=ro", uri=True) as con:
        return con.execute(sql).fetchall()


def money(value: float) -> str:
    return f"S${value:,.2f}"


def draft(
    api: Api, llm: ScriptedLLMClient, question: str, *turns: Any, answer: str = "Your draft is ready."
) -> Any:
    route(llm, question, "drafting")
    llm.script_turns("drafting", question, *turns)
    llm.script("merge", question, answer)
    return api.run(api.create_thread(), question)


def submit(template_id: str, slots: dict[str, str], unfilled: list[dict[str, str]] | None = None) -> Any:
    return turn(tool_call("submit_draft", template_id=template_id, slots=slots, unfilled=unfilled or []))


def tool_results(llm: ScriptedLLMClient) -> list[str]:
    last = max(llm.sent, key=len)
    return [
        str(b["content"]) for m in last if m["role"] == "user" and isinstance(m["content"], list)
        for b in m["content"] if b["type"] == "tool_result"
    ]


def reminder_slots(**overrides: str) -> dict[str, str]:
    rows = records(INVOICES)
    total = sum(r[3] for r in rows)
    table = "\n".join(f"{r[0]} | {r[1]} | due {r[2]} | {money(r[3])}" for r in rows)
    slots = {
        "letter_date": "30 September 2026",
        "tenant_name": "Golden Leaf Minimart Pte Ltd",
        "tenant_address": "#01-32, Blk 214 Bayshore Vale Drive, Singapore 460214",
        "unit_address": "Blk 214 Bayshore Vale Drive #01-32",
        "tenancy_id": "TA-2024-00542",
        "notice_level": "FIRST",
        "invoice_table": table,
        "total_due": money(total),
        "pay_by_date": "14 October 2026",
        "officer_name": "Siti Aishah Yusof",
        "officer_phone": "+65 6100 2058",
    }
    return {**slots, **overrides}


def test_arrears_reminder_is_filled_from_the_records(api: Api, llm: ScriptedLLMClient) -> None:
    rows = records(INVOICES)
    assert len(rows) >= 2
    total = money(sum(r[3] for r in rows))

    stream = draft(
        api, llm, REMINDER,
        turn(tool_call("lookup_records", plan="Unpaid invoices of the tenancy", sql=INVOICES)),
        submit("TPL-REM-01", reminder_slots(), [{"slot": "letter_ref", "reason": "No letter reference was given."}]),
        turn(text("Drafted the first reminder.")),
    )

    (event,) = stream.of_type("draft")
    assert (event["template_id"], event["template_version"]) == ("TPL-REM-01", "1.3")
    html = str(event["html"])
    assert total in html
    assert all(r[0] in html for r in rows)
    assert "Golden Leaf Minimart Pte Ltd" in html
    assert "FIRST REMINDER: OVERDUE RENT" in html
    assert "[Officer:" not in html and "{{" not in html  # no guidance text, no raw slots
    assert "Approved wording" not in html  # the template's document-control tables are not part of the letter
    assert any(r[0] in "\n".join(tool_results(llm)) for r in rows)  # the lookup ran against real records
    assert stream.of_type("done") and not stream.of_type("error")
    assert [e["node"] for e in stream.of_type("node") if e["status"] == "started"] == ["route", "drafting", "merge"]


def test_slots_that_cannot_be_filled_are_listed(api: Api, llm: ScriptedLLMClient) -> None:
    slots = reminder_slots()
    del slots["officer_phone"]  # the agent never even mentioned it

    stream = draft(
        api, llm, REMINDER,
        submit("TPL-REM-01", slots, [{"slot": "letter_ref", "reason": "No letter reference was given."}]),
        turn(text("Drafted.")),
    )

    (event,) = stream.of_type("draft")
    unfilled = {u["slot"]: u["reason"] for u in event["unfilled"]}
    assert set(unfilled) == {"letter_ref", "officer_phone"}
    assert unfilled["letter_ref"] == "No letter reference was given."
    assert unfilled["officer_phone"]  # a reason is always given
    assert "letter_ref" in str(event["html"])  # visibly marked in the preview, not silently blank
    assert event["pii"] == []


def test_missing_documents_notice_lists_what_the_checklist_says_is_missing(api: Api, llm: ScriptedLLMClient) -> None:
    missing = records(
        "SELECT doc_label FROM application_documents WHERE application_id = 'APP-2025-0012' AND mandatory = 1 AND submitted = 0"
    )
    assert len(missing) >= 2
    items = "\n".join(f"{i}. {label} (missing)" for i, (label,) in enumerate(missing, start=1))

    stream = draft(
        api, llm, NOTICE,
        turn(tool_call("get_template", template_id="TPL-DOC-02")),
        submit("TPL-DOC-02", {
            "letter_date": "30 September 2026", "tenant_name": "Example Pte Ltd", "contact_person": "Ms Tan",
            "application_id": "APP-2025-0012", "application_type": "NEW TENANCY", "unit_address": "Blk 1 #01-01",
            "missing_documents_list": items, "submit_by_date": "14 October 2026",
            "officer_name": "Siti Aishah Yusof", "officer_email": "siti.aishah@heb.example.org",
        }, [{"slot": "letter_ref", "reason": "No letter reference was given."}]),
        turn(text("Drafted.")),
    )

    (event,) = stream.of_type("draft")
    assert event["template_id"] == "TPL-DOC-02"
    assert all(label in str(event["html"]) for (label,) in missing)
    assert "Please submit the documents by {{ submit_by_date }}" in "\n".join(tool_results(llm))  # the agent was shown the template


def test_personal_data_in_a_draft_is_flagged(api: Api, llm: ScriptedLLMClient) -> None:
    contact = records("SELECT contact_person, phone FROM tenants WHERE name = 'Golden Leaf Minimart Pte Ltd'")[0]

    stream = draft(
        api, llm, REPLY,
        submit("TPL-FBK-03", {
            "feedback_id": "FB-2025-00002", "reply_date": "30 September 2026", "provider_name": "Mr Lim",
            "received_date": "2025-10-02", "location": "Blk 675 #02-07",
            "issue_summary": "a leaking ceiling",
            "action_taken": f"We spoke to the tenant, {contact[0]}, whose NRIC is S1234567D.",
            "ticket_reference": "IEMS-2510-0002",
            "next_steps": f"You may call the tenant on {contact[1]}, and write to me at siti.aishah@heb.example.org.",
            "officer_name": "Siti Aishah Yusof",
        }),
        turn(text("Drafted.")),
    )

    (event,) = stream.of_type("draft")
    flagged = {(p["kind"], p["slot"]) for p in event["pii"]}
    assert ("nric", "action_taken") in flagged
    assert ("phone", "next_steps") in flagged
    assert ("tenant", "action_taken") in flagged  # the tenant's name in a reply to a member of the public
    assert ("email", "next_steps") not in flagged  # the officer's own contact details are expected
    assert "S1234567D" not in str(event["pii"])  # findings point at the slot; they don't repeat the number


def test_the_officers_own_contact_details_are_not_flagged(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots()), turn(text("Drafted.")))

    (event,) = stream.of_type("draft")
    assert event["pii"] == []


def test_the_draft_downloads_as_docx(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots()), turn(text("Drafted.")))
    (event,) = stream.of_type("draft")

    response = api.client.get(f"/drafts/{event['draft_id']}/file", headers=USER)

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX
    assert "attachment" in response.headers["content-disposition"] and ".docx" in response.headers["content-disposition"]
    body = "\n".join(p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs)
    assert "Golden Leaf Minimart Pte Ltd" in body
    assert "TPL-REM-01" not in body and "Approved wording" not in body and "{{" not in body


def test_a_draft_is_only_downloadable_by_the_officer_who_asked(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots()), turn(text("Drafted.")))
    (event,) = stream.of_type("draft")
    path = f"/drafts/{event['draft_id']}/file"

    assert api.client.get(path, headers={"X-PLG-User": "S04"}).status_code == 404
    assert api.client.get(path).status_code == 401
    assert api.client.get("/drafts/nope/file", headers=USER).status_code == 404


def test_an_unknown_template_is_rejected_and_no_draft_is_made(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(
        api, llm, REMINDER,
        submit("TPL-XXX-99", {"tenant_name": "x"}),
        turn(text("There is no approved template for that.")),
        answer="I couldn't find an approved template for that.",
    )

    assert not stream.of_type("draft")
    assert "TPL-REM-01" in tool_results(llm)[0]  # the agent is told which templates exist
    assert stream.of_type("done") and not stream.of_type("error")


def test_drafting_agent_can_only_read_and_propose(api: Api, llm: ScriptedLLMClient) -> None:
    draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots()), turn(text("Drafted.")))

    _, tools = llm.requests[1]  # request 0 is the router
    assert {t["name"] for t in tools} == {"get_template", "lookup_records", "submit_draft"}


def test_the_agent_is_offered_only_approved_templates(api: Api, llm: ScriptedLLMClient) -> None:
    draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots()), turn(text("Drafted.")))

    system, _ = llm.requests[1]
    prompt = "\n".join(str(b["text"]) for b in system)
    for template_id in ("TPL-REM-01", "TPL-DOC-02", "TPL-FBK-03"):
        assert template_id in prompt
    assert "Siti Aishah Yusof" in prompt and "siti.aishah@heb.example.org" in prompt  # who the letters are signed by


def test_a_later_draft_replaces_an_earlier_one(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(
        api, llm, REMINDER,
        submit("TPL-REM-01", reminder_slots(notice_level="FIRST")),
        submit("TPL-REM-01", reminder_slots(notice_level="SECOND")),
        turn(text("Drafted.")),
    )

    (event,) = stream.of_type("draft")
    assert "SECOND REMINDER" in str(event["html"])
    assert api.client.get(f"/drafts/{event['draft_id']}/file", headers=USER).status_code == 200


def test_a_value_that_looks_like_an_open_slot_is_not_marked_open(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots(tenant_name="«fake» Pte Ltd")), turn(text("x")))

    (event,) = stream.of_type("draft")
    assert "<mark>«fake»" not in str(event["html"])


def test_open_slots_are_highlighted_in_the_docx(api: Api, llm: ScriptedLLMClient) -> None:
    stream = draft(api, llm, REMINDER, submit("TPL-REM-01", reminder_slots(officer_phone="")), turn(text("x")))
    (event,) = stream.of_type("draft")

    response = api.client.get(f"/drafts/{event['draft_id']}/file", headers=USER)

    runs = [r for p in docx.Document(io.BytesIO(response.content)).paragraphs for r in p.runs if "«officer_phone»" in r.text]
    assert runs and all(r.font.highlight_color is not None for r in runs)
